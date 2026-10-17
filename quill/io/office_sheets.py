"""Spreadsheets as readable text without MarkItDown (.xlsx, .xlsm, .ods).

Part of QUILL's own Office readers (:mod:`quill.io.office_text`). Each sheet
becomes a ``## Sheet: <name>`` heading followed by a Markdown table, so a
screen-reader user moves sheet to sheet with heading navigation and cell to
cell with QUILL's table commands. The first row is the table's header row,
because that is what it is in almost every real workbook.

Two readers, tried in order for ``.xlsx``:

1. **openpyxl** (shipped in the ``office-text`` extra): cached values, so a
   formula shows its last result, and dates come back as dates.
2. **The workbook's own XML** (standard library only): shared strings, inline
   strings, numbers, booleans, and dates (a date is stored as a day count;
   the cell's number format in ``styles.xml`` is what says it is a date).
   Used when openpyxl is missing or cannot open the file. A formula shows its
   last calculated result, as in openpyxl; a formula never calculated shows
   nothing.

``.ods`` (OpenDocument) is read from ``content.xml`` the same way.

Reads are bounded (PERF-11): at most :data:`MAX_ROWS` rows and
:data:`MAX_COLS` columns per sheet, and a sheet that was cut short says so in
the result's notes, which the caller reports instead of hiding.

wx-free; strict-typed.
"""

from __future__ import annotations

import posixpath
import re
from dataclasses import dataclass
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any
from xml.etree.ElementTree import Element

from quill.core.safe_archive import open_zip
from quill.core.safe_xml import fromstring as safe_xml_fromstring

__all__ = [
    "MAX_COLS",
    "MAX_ROWS",
    "SheetText",
    "escape_cell",
    "markdown_table",
    "read_ods_text",
    "read_xlsx_text",
    "render_workbook",
]

#: Per-sheet bounds. A workbook bigger than this opens with what fits and a
#: note naming the sheet that was cut short -- never silently, never by
#: exhausting memory on a sheet with a million rows.
MAX_ROWS = 2000
MAX_COLS = 50

_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_PKG_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
_ODF_TABLE = "urn:oasis:names:tc:opendocument:xmlns:table:1.0"
_ODF_TEXT = "urn:oasis:names:tc:opendocument:xmlns:text:1.0"
_ODF_OFFICE = "urn:oasis:names:tc:opendocument:xmlns:office:1.0"
#: Day 0 of Excel's 1900 date system, as Excel counts it (its 1900 leap-year
#: bug folded in, so every date from March 1900 on is right).
_EPOCH = datetime(1899, 12, 30)  # noqa: DTZ001 - a calendar date, not a moment


@dataclass(frozen=True, slots=True)
class SheetText:
    """A workbook rendered as text, and how it was read."""

    text: str
    engine: str
    sheets: int
    notes: tuple[str, ...] = ()


def escape_cell(text: str) -> str:
    """One cell flattened to a line and escaped for a Markdown table.

    A ``|`` inside a cell would otherwise split it into two columns and throw
    off "column 2 of 4" in table navigation; GFM's escape keeps it cell text.
    """
    flat = " ".join(str(text).split())
    return flat.replace("\\", "\\\\").replace("|", "\\|")


def markdown_table(rows: list[list[str]]) -> list[str]:
    """Rows of already-escaped cells as a GFM table; ``[]`` when all are empty.

    Wholly empty rows are dropped (a spreadsheet's blank spacer rows would
    otherwise read as rows of nothing), and short rows are padded so every row
    has the same number of cells.
    """
    kept = [row for row in rows if any(cell.strip() for cell in row)]
    if not kept:
        return []
    width = max(len(row) for row in kept)
    padded = [row + [""] * (width - len(row)) for row in kept]
    lines = ["| " + " | ".join(padded[0]) + " |", "| " + " | ".join(["---"] * width) + " |"]
    lines.extend("| " + " | ".join(row) + " |" for row in padded[1:])
    return lines


def _trim(row: list[str]) -> list[str]:
    while row and not row[-1].strip():
        row.pop()
    return row


def _sheet_lines(name: str, rows: list[list[str]], truncated: bool, notes: list[str]) -> list[str]:
    lines = [f"## Sheet: {name}", ""]
    lines.extend(markdown_table(rows) or ["(This sheet is empty.)"])
    lines.append("")
    if truncated:
        notes.append(
            f'Sheet "{name}" is larger than {MAX_ROWS} rows or {MAX_COLS} columns; '
            "only that much of it was read."
        )
    return lines


def render_workbook(workbook: Any, notes: list[str] | None = None) -> str:
    """An openpyxl workbook (or anything shaped like one) as sheet-headed tables.

    Rows are pulled lazily and the iterator is abandoned at :data:`MAX_ROWS`,
    so a huge read-only sheet is never materialised.
    """
    found = notes if notes is not None else []
    lines: list[str] = []
    for worksheet in workbook.worksheets:
        rows: list[list[str]] = []
        truncated = False
        for row in worksheet.iter_rows(values_only=True):
            cells = ["" if cell is None else escape_cell(_value_text(cell)) for cell in row]
            truncated = truncated or len(cells) > MAX_COLS
            rows.append(_trim(cells[:MAX_COLS]))
            if len(rows) >= MAX_ROWS:
                total = getattr(worksheet, "max_row", None)
                truncated = truncated or (isinstance(total, int) and total > MAX_ROWS)
                break
        lines.extend(_sheet_lines(str(worksheet.title), rows, truncated, found))
    return "\n".join(lines).rstrip() + "\n"


def _value_text(value: object) -> str:
    """A cell value as a person would read it: whole numbers without ``.0``."""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    isoformat = getattr(value, "isoformat", None)
    if callable(isoformat):
        text = str(isoformat())
        return text[:-9] if text.endswith("T00:00:00") else text.replace("T", " ")
    return str(value)


# ---------------------------------------------------------------- .xlsx


def read_xlsx_text(path: Path) -> SheetText:
    """Read an ``.xlsx``/``.xlsm``: openpyxl first, then the workbook's own XML.

    Raises ``ValueError`` (with a sentence fit to say aloud) when neither can
    read it, and lets ``zipfile.BadZipFile`` through for a file that is not a
    workbook at all, so the caller can say "damaged" rather than "unsupported".
    """
    notes: list[str] = []
    try:
        from openpyxl import load_workbook  # type: ignore[import-untyped]
    except ImportError:
        load_workbook = None
    if load_workbook is not None:
        try:
            workbook = load_workbook(str(path), data_only=True, read_only=True)
        except Exception:  # noqa: BLE001 - any openpyxl failure tries the XML reader
            workbook = None
        if workbook is not None:
            try:
                text = render_workbook(workbook, notes)
                return SheetText(text, "openpyxl", len(workbook.worksheets), tuple(notes))
            except Exception:  # noqa: BLE001 - a half-read workbook tries the XML reader
                notes.clear()
            finally:
                workbook.close()
    return _read_xlsx_xml(path)


def _read_xlsx_xml(path: Path) -> SheetText:
    with open_zip(path) as archive:
        names = set(archive.namelist())
        if "xl/workbook.xml" not in names:
            raise ValueError(f"{path.name} has no workbook inside it.")
        shared = (
            _shared_strings(archive.read("xl/sharedStrings.xml"))
            if ("xl/sharedStrings.xml" in names)
            else []
        )
        targets = (
            _workbook_targets(archive.read("xl/_rels/workbook.xml.rels"))
            if ("xl/_rels/workbook.xml.rels" in names)
            else {}
        )
        dates = _date_styles(archive.read("xl/styles.xml")) if "xl/styles.xml" in names else set()
        workbook = safe_xml_fromstring(archive.read("xl/workbook.xml"))
        notes: list[str] = []
        lines: list[str] = []
        count = 0
        for sheet in workbook.iter(f"{{{_MAIN}}}sheet"):
            name = sheet.get("name") or f"Sheet {count + 1}"
            target = targets.get(sheet.get(f"{{{_REL}}}id") or "", "")
            if target not in names:
                continue
            sheet_root = safe_xml_fromstring(archive.read(target))
            rows, truncated = _xlsx_rows(sheet_root, shared, dates)
            lines.extend(_sheet_lines(name, rows, truncated, notes))
            count += 1
    if not count:
        raise ValueError(f"{path.name} has no sheets QUILL could read.")
    return SheetText("\n".join(lines).rstrip() + "\n", "sheet-xml", count, tuple(notes))


def _shared_strings(data: bytes) -> list[str]:
    root = safe_xml_fromstring(data)
    found: list[str] = []
    for item in root.findall(f"{{{_MAIN}}}si"):
        # Plain text, or rich-text runs; never the phonetic guide (rPh) a
        # Japanese workbook stores beside the text it annotates.
        plain = item.find(f"{{{_MAIN}}}t")
        runs = item.findall(f"{{{_MAIN}}}r/{{{_MAIN}}}t")
        found.append(plain.text or "" if plain is not None else "".join(r.text or "" for r in runs))
    return found


#: Excel's built-in number formats that show a date or a time.
_BUILTIN_DATE_FORMATS = frozenset({*range(14, 23), 45, 46, 47})
_QUOTED = re.compile(r'"[^"]*"|\[[^\]]*\]|\\.')


def _date_styles(data: bytes) -> set[int]:
    """Indexes of the cell styles (``s="N"``) that display a date or time.

    A date is stored as a day count; only its style says it is a date. A
    custom format counts when, outside quoted text and ``[...]`` codes, it
    has a day, month, year, hour or second placeholder.
    """
    root = safe_xml_fromstring(data)
    custom: dict[int, bool] = {}
    for fmt in root.iter(f"{{{_MAIN}}}numFmt"):
        code = _QUOTED.sub("", (fmt.get("formatCode") or "").lower())
        custom[int(fmt.get("numFmtId") or -1)] = any(ch in code for ch in "dmyhs")
    styles: set[int] = set()
    cell_formats = root.find(f"{{{_MAIN}}}cellXfs")
    for index, xf in enumerate(cell_formats if cell_formats is not None else []):
        number_format = int(xf.get("numFmtId") or 0)
        if number_format in _BUILTIN_DATE_FORMATS or custom.get(number_format, False):
            styles.add(index)
    return styles


def _serial_date(serial: float) -> str:
    """An Excel day count as an ISO date (with the time when it has one)."""
    moment = _EPOCH + timedelta(days=serial)
    if moment.hour == moment.minute == moment.second == 0:
        return moment.date().isoformat()
    return moment.isoformat(sep=" ", timespec="minutes")


def _workbook_targets(data: bytes) -> dict[str, str]:
    root = safe_xml_fromstring(data)
    targets: dict[str, str] = {}
    for relation in root.findall(f"{{{_PKG_REL}}}Relationship"):
        target = relation.get("Target") or ""
        resolved = (
            target.lstrip("/")
            if target.startswith("/")
            else posixpath.normpath(posixpath.join("xl", target))
        )
        targets[relation.get("Id") or ""] = resolved
    return targets


_COLUMN = re.compile(r"^([A-Z]+)")


def _column_index(reference: str) -> int:
    match = _COLUMN.match(reference or "")
    if match is None:
        return -1
    index = 0
    for letter in match.group(1):
        index = index * 26 + (ord(letter) - 64)
    return index - 1


def _xlsx_rows(sheet: Element, shared: list[str], dates: set[int]) -> tuple[list[list[str]], bool]:
    rows: list[list[str]] = []
    truncated = False
    for row in sheet.iter(f"{{{_MAIN}}}row"):
        if len(rows) >= MAX_ROWS:
            truncated = True
            break
        cells: list[str] = []
        for cell in row.findall(f"{{{_MAIN}}}c"):
            index = _column_index(cell.get("r") or "")
            index = len(cells) if index < 0 else index
            if index >= MAX_COLS:
                truncated = True
                continue
            cells.extend([""] * (index - len(cells)))
            cells.append(escape_cell(_xlsx_value(cell, shared, dates)))
        rows.append(_trim(cells))
    return rows, truncated


def _xlsx_value(cell: Element, shared: list[str], dates: set[int]) -> str:
    kind = cell.get("t") or "n"
    if kind == "inlineStr":
        return "".join(node.text or "" for node in cell.iter(f"{{{_MAIN}}}t"))
    raw = cell.findtext(f"{{{_MAIN}}}v") or ""
    if kind == "s":
        try:
            return shared[int(raw)]
        except (ValueError, IndexError):
            return ""
    if kind == "b":
        return "TRUE" if raw == "1" else "FALSE"
    if kind == "n" and raw:
        try:
            number = float(raw)
        except ValueError:
            return raw
        style = cell.get("s") or ""
        if style.isdigit() and int(style) in dates and 0 <= number < 2958466:
            return _serial_date(number)
        return _value_text(number)
    return raw


# ---------------------------------------------------------------- .ods


def read_ods_text(path: Path) -> SheetText:
    """Read an OpenDocument spreadsheet (``.ods``) from its ``content.xml``."""
    with open_zip(path) as archive:
        root = safe_xml_fromstring(archive.read("content.xml"))
    notes: list[str] = []
    lines: list[str] = []
    count = 0
    for spreadsheet in root.iter(f"{{{_ODF_OFFICE}}}spreadsheet"):
        for table in spreadsheet.findall(f"{{{_ODF_TABLE}}}table"):
            count += 1
            name = table.get(f"{{{_ODF_TABLE}}}name") or f"Sheet {count}"
            rows, truncated = _ods_rows(table)
            lines.extend(_sheet_lines(name, rows, truncated, notes))
    if not count:
        raise ValueError(f"{path.name} has no sheets QUILL could read.")
    return SheetText("\n".join(lines).rstrip() + "\n", "odf-xml", count, tuple(notes))


def _repeat(node: Element, attribute: str) -> int:
    try:
        return max(1, int(node.get(f"{{{_ODF_TABLE}}}{attribute}") or "1"))
    except ValueError:
        return 1


def _ods_rows(table: Element) -> tuple[list[list[str]], bool]:
    rows: list[list[str]] = []
    truncated = False
    for row in table.iter(f"{{{_ODF_TABLE}}}table-row"):
        cells: list[str] = []
        for cell in row:
            if not cell.tag.endswith(("}table-cell", "}covered-table-cell")):
                continue
            text = escape_cell(
                " ".join(
                    "".join(paragraph.itertext()) for paragraph in cell.iter(f"{{{_ODF_TEXT}}}p")
                )
            )
            # An empty cell repeated 16,000 times is how ODF says "the rest of
            # the row is blank"; only a repeat with something in it is real.
            repeat = _repeat(cell, "number-columns-repeated") if text else 1
            cells.extend([text] * min(repeat, MAX_COLS + 1 - len(cells)))
            if len(cells) > MAX_COLS:
                truncated = truncated or any(cells[MAX_COLS:])
                cells = cells[:MAX_COLS]
                break
        cells = _trim(cells)
        repeat = _repeat(row, "number-rows-repeated") if cells else 1
        for _ in range(repeat):
            if len(rows) >= MAX_ROWS:
                return rows, True
            rows.append(list(cells))
    return rows, truncated
