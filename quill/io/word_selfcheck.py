"""Prove, inside a built bundle, that Word documents really open and save.

Why this exists (2026-10-17): the shared QuillVille runtime excluded ``lxml``
for months. python-docx is pure Python and froze in without complaint, so
``find_spec("docx")`` answered yes -- and then ``import docx`` raised the
moment it reached for ``lxml``. Every installed copy of QUILL and QUILL Lite
quietly fell back to opening Word files as plain text, could not save a real
``.docx``, and the GLOW Word, PDF and PowerPoint audits failed. Every build
exited 0, because nothing asked the bundle to do the actual job.

This module is that question, asked the same way everywhere a bundle is built:

* ``scripts/check_runtime_imports.py`` runs it inside the frozen shared runtime;
* ``standalone/studio/scripts/build_portable.py`` runs it inside the QUILL Lite
  portable (and any product whose ``must_import`` names this module);
* ``scripts/build_windows_distribution.py`` runs it inside QUILL's embedded
  runtime;
* ``tests/unit/scripts/test_word_support_in_runtime.py`` runs it with every
  package the runtime spec still excludes blocked, so a new exclude that breaks
  Word fails a pull request rather than a release.

It imports the Word stack, writes a small real ``.docx`` with python-docx (a
heading, a bold word, a two-by-two table), opens it the way both editors do
(``quill.core.docx_open.read_docx_as_rtf``), saves the result back to Word
bytes (``rtf_to_docx_bytes``), and reads those bytes again. Standard library
plus the modules it is checking; no ``wx``.
"""

from __future__ import annotations

import importlib
import io
import tempfile
from pathlib import Path

__all__ = ["WORD_SUPPORT_MODULES", "assert_word_round_trip", "word_round_trip_problem"]

#: What opening and saving a Word document needs, in the order a failure is
#: most useful to read: the native parser first, then python-docx on top of it,
#: then QUILL's own reader, writer and the shared open/save decisions.
WORD_SUPPORT_MODULES: tuple[str, ...] = (
    "lxml.etree",
    "docx",
    "quill.io.docx_reader",
    "quill.io.docx_writer",
    "quill.io.docx_text",
    "quill.core.docx_open",
)

_HEADING = "Word check heading"
_BOLD = "boldword"
_CELL = "cell-b2"


def _sample_docx() -> bytes:
    import docx

    document = docx.Document()
    document.add_heading(_HEADING, level=1)
    paragraph = document.add_paragraph("Plain text and a ")
    paragraph.add_run(_BOLD).bold = True
    table = document.add_table(rows=2, cols=2)
    table.cell(1, 1).text = _CELL
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def word_round_trip_problem() -> str | None:
    """``None`` when Word works in this interpreter; otherwise what broke, in a sentence."""
    for name in WORD_SUPPORT_MODULES:
        try:
            importlib.import_module(name)
        except Exception as error:  # noqa: BLE001 - reported, never raised
            return f"cannot import {name}: {type(error).__name__}: {error}"
    try:
        import docx

        from quill.core.docx_open import (
            fidelity_findings,
            read_docx_as_rtf,
            rtf_to_docx_bytes,
            word_rich_available,
        )

        if not word_rich_available():
            return "word_rich_available() is False, so both editors would open Word as text"
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "word-check.docx"
            source.write_bytes(_sample_docx())
            if not any(f.startswith("tables") for f in fidelity_findings(source)):
                return "the fidelity scan did not notice the table in a sample document"
            rtf = read_docx_as_rtf(source)
            for expected in (_HEADING, _BOLD, _CELL):
                if expected not in rtf:
                    return f"opening the sample document lost {expected!r}"
            saved = docx.Document(io.BytesIO(rtf_to_docx_bytes(rtf)))
        texts = [p.text for p in saved.paragraphs]
        if not any(_HEADING in text for text in texts):
            return "saving back to Word lost the heading"
        heading = next(p for p in saved.paragraphs if _HEADING in p.text)
        style_name = str(getattr(heading.style, "name", "") or "")
        if not style_name.startswith("Heading"):
            return f"saving back to Word turned the heading into {style_name!r}"
        if not any(r.bold and _BOLD in r.text for p in saved.paragraphs for r in p.runs):
            return "saving back to Word lost the bold word"
    except Exception as error:  # noqa: BLE001 - reported, never raised
        return f"the Word round trip raised {type(error).__name__}: {error}"
    return None


def assert_word_round_trip() -> None:
    """Raise ``RuntimeError`` naming what broke; the build scripts call this."""
    problem = word_round_trip_problem()
    if problem is not None:
        raise RuntimeError(
            "Word documents would not open or save as Word in this build: "
            f"{problem}. python-docx needs lxml; check that no spec or "
            "build step excludes or prunes lxml or docx."
        )
