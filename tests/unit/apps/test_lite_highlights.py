"""Highlighted passages in QUILL Lite: QUILL's three commands, shared, on QUILL's chords.

Every handler is *called* here (GATE-LITE-COVER). QUILL Lite opens a Word file
as read-only text, so its highlights are read from the file itself; an RTF
document's come from the rich control.
"""

from __future__ import annotations

from pathlib import Path

import pytest

from quill.ui.highlight_commands import NO_HIGHLIGHTS


def _word_file(tmp_path: Path) -> Path:
    docx = pytest.importorskip("docx")
    from docx.enum.text import WD_COLOR_INDEX

    document = docx.Document()
    paragraph = document.add_paragraph("Pay the ")
    paragraph.add_run("deposit").font.highlight_color = WD_COLOR_INDEX.YELLOW
    paragraph.add_run(" by Friday.")
    paragraph = document.add_paragraph("Bring ")
    paragraph.add_run("two forms of ID").font.highlight_color = WD_COLOR_INDEX.PINK
    path = tmp_path / "lease.docx"
    document.save(str(path))
    return path


def _opened_word_file(lite_window, tmp_path: Path):
    from quill.io.office_text import read_office_text_preferring_own

    path = _word_file(tmp_path)
    text = read_office_text_preferring_own(path).text
    win = lite_window(text, cursor=0)
    win.path = path
    return win, text


def test_next_and_previous_walk_a_word_files_highlights(lite_window, tmp_path) -> None:
    win, text = _opened_word_file(lite_window, tmp_path)
    win.cmd_next_highlight()
    assert text[win.control.GetInsertionPoint() :].startswith("deposit")
    assert win.announcements[-1] == "Yellow highlight, 1 of 2: deposit"
    win.cmd_next_highlight()
    assert win.announcements[-1] == "Pink highlight, 2 of 2: two forms of ID"
    win.cmd_previous_highlight()
    assert win.announcements[-1] == "Yellow highlight, 1 of 2: deposit"


def test_the_list_lands_on_the_chosen_passage(lite_window, tmp_path, monkeypatch) -> None:
    win, text = _opened_word_file(lite_window, tmp_path)
    rows: list[str] = []

    def fake(_wx, _parent, _show, passages, start):
        rows.extend(p.text for p in passages)
        return 1

    monkeypatch.setattr("quill.ui.highlight_list_dialog.show_highlight_list", fake)
    win.cmd_list_highlights()
    assert rows == ["deposit", "two forms of ID"]
    assert text[win.control.GetInsertionPoint() :].startswith("two forms")


def test_a_rich_document_reads_the_control(lite_window) -> None:
    from quill.io.rtf import markdown_to_rtf

    win = lite_window("Keep this. Mark that.", cursor=0, mode="rich")
    win.editor.rtf_bytes = markdown_to_rtf('Keep this. [Mark that]{highlight="green"}.').encode(
        "utf-8"
    )
    win.cmd_next_highlight()
    assert win.announcements[-1] == "Green highlight, 1 of 1: Mark that"
    assert win.control.GetInsertionPoint() == len("Keep this. ")


def test_no_highlights_is_said_plainly(lite_window) -> None:
    win = lite_window("Nothing marked here.", cursor=0)
    for command in (
        lambda w: w.cmd_next_highlight(),
        lambda w: w.cmd_previous_highlight(),
        lambda w: w.cmd_list_highlights(),
    ):
        command(win)
        assert win.announcements[-1] == NO_HIGHLIGHTS
