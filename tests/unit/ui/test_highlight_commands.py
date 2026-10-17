"""Highlighted passages in QUILL: the shared commands against QUILL's host.

Item 39 (2026-10-16). Three sources, each tested where it is real:

* a **rich** document in a real Rich Edit control (built in memory, never
  shown, no input sent): imported highlights, one applied with Format >
  Highlight's own method, one removed with Clear Formatting's, and an edit
  above a passage;
* a **Markdown** document's hidden codes, through QUILL's own host hooks;
* a **Word file opened as text**, whose highlights come from the file.

The list dialog is driven for real once (Escape and Enter), and replaced by a
recorder elsewhere.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

import pytest

from quill.ui.highlight_commands import NO_HIGHLIGHTS, HighlightCommandsMixin
from quill.ui.main_frame_highlights import HighlightPassagesMixin
from tests.wx_session import shared_wx_app

wx = pytest.importorskip("wx")


class _Control:
    def __init__(self, text: str, caret: int = 0) -> None:
        self.text = text
        self.caret = caret

    def GetValue(self) -> str:  # noqa: N802 - wx spelling
        return self.text

    def GetInsertionPoint(self) -> int:  # noqa: N802
        return self.caret

    def SetInsertionPoint(self, position: int) -> None:  # noqa: N802
        self.caret = position


class _Host(HighlightCommandsMixin):
    """The shared mixin with QUILL Lite's default hooks, recorded."""

    def __init__(self, text: str, *, kind: str = "markdown", path: Path | None = None) -> None:
        self.control = _Control(text)
        self.kind = kind
        self.path = path
        self.said: list[str] = []
        self.went: list[int] = []

    def _host_kind(self) -> str:
        return self.kind

    def _host_say(self, message: str) -> None:
        self.said.append(message)

    def _go_to(self, position: int) -> None:
        self.went.append(position)
        self.control.caret = position

    def _host_parent(self) -> Any:
        return None


MD = (
    "# Contract\n\n"
    'The [budget is final]{highlight="yellow"} unless the [board]{highlight="green"} objects.\n'
)


def test_next_and_previous_walk_markdown_highlights_and_wrap() -> None:
    host = _Host(MD)
    host.cmd_next_highlight()
    assert MD[host.control.caret :].startswith("budget is final")
    assert host.said[-1] == "Yellow highlight, 1 of 2: budget is final"
    host.cmd_next_highlight()
    assert host.said[-1] == "Green highlight, 2 of 2: board"
    host.cmd_next_highlight()
    assert host.said[-1].startswith("Yellow highlight, 1 of 2"), "wraps"
    host.cmd_previous_highlight()
    assert host.said[-1] == "Green highlight, 2 of 2: board"


def test_nothing_highlighted_is_said_plainly() -> None:
    host = _Host("plain words only")
    for command in (
        lambda h: h.cmd_next_highlight(),
        lambda h: h.cmd_previous_highlight(),
        lambda h: h.cmd_list_highlights(),
    ):
        command(host)
        assert host.said[-1] == NO_HIGHLIGHTS


def test_the_list_returns_to_the_chosen_passage(monkeypatch) -> None:
    seen: dict[str, Any] = {}

    def fake(_wx, _parent, _show, passages, start):
        seen["rows"] = [p.text for p in passages]
        seen["start"] = start
        return 1

    monkeypatch.setattr("quill.ui.highlight_list_dialog.show_highlight_list", fake)
    host = _Host(MD)
    host.control.caret = MD.index("unless")
    host.cmd_list_highlights()
    assert seen == {"rows": ["budget is final", "board"], "start": 1}
    assert MD[host.control.caret :].startswith("board")
    assert host.said == [], "arrival from the list is the reader's to read"


def test_escape_from_the_list_leaves_the_caret(monkeypatch) -> None:
    monkeypatch.setattr(
        "quill.ui.highlight_list_dialog.show_highlight_list", lambda *_a, **_k: None
    )
    host = _Host(MD)
    host.control.caret = 3
    host.cmd_list_highlights()
    assert host.control.caret == 3 and host.went == []


def test_a_word_file_shown_as_text_reads_its_highlights_from_the_file(tmp_path: Path) -> None:
    docx = pytest.importorskip("docx")
    from docx.enum.text import WD_COLOR_INDEX

    document = docx.Document()
    paragraph = document.add_paragraph("Sign ")
    paragraph.add_run("here").font.highlight_color = WD_COLOR_INDEX.YELLOW
    paragraph.add_run(" and date it.")
    path = tmp_path / "form.docx"
    document.save(str(path))
    # What read-extract shows, with a line typed above it since opening.
    host = _Host("A note I added.\nSign here and date it.\n", kind="plain", path=path)
    host.cmd_next_highlight()
    assert host.control.text[host.control.caret :].startswith("here and date")
    assert host.said[-1] == "Yellow highlight, 1 of 1: here"


# --------------------------------------------------------------------------- #
# A real Rich Edit control
# --------------------------------------------------------------------------- #


@pytest.fixture
def rich_control():
    from quill.ui.richedit_rtf_surface import create_richedit_rtf

    app = shared_wx_app()
    frame = wx.Frame(None)
    control = create_richedit_rtf(wx, frame, wx.TE_MULTILINE)
    wrapper = getattr(control, "quill_richedit", None)
    if wrapper is None or not wrapper.rtf_available():
        frame.Destroy()
        pytest.skip("no Text Object Model on this machine")
    yield control, wrapper
    frame.Destroy()
    del app


class _RichHost(_Host):
    def __init__(self, control: Any, wrapper: Any) -> None:
        super().__init__("")
        self.control = control
        self.wrapper = wrapper

    def _host_rich_wrapper(self) -> Any:
        return self.wrapper

    def _go_to(self, position: int) -> None:
        self.went.append(position)
        self.control.SetInsertionPoint(position)


def _load_word_shape(wrapper: Any) -> None:
    from quill.io.rtf_model import InlineSpan, RichDocument, RichParagraph, rich_to_rtf

    document = RichDocument([
        RichParagraph([InlineSpan("Terms", bold=True)], style="heading", level=1),
        RichParagraph([
            InlineSpan("The "),
            InlineSpan("deposit", highlight="yellow"),
            InlineSpan(" is due before the "),
            InlineSpan("move-in date", highlight="pink"),
            InlineSpan("."),
        ]),
    ])
    wrapper.set_rtf(rich_to_rtf(document).encode("utf-8"))


def test_rich_imported_applied_and_removed_highlights(rich_control) -> None:
    control, wrapper = rich_control
    _load_word_shape(wrapper)
    host = _RichHost(control, wrapper)
    passages = host._highlight_passages()
    text = control.GetValue()
    assert [(p.colour, text[p.start : p.end]) for p in passages] == [
        ("yellow", "deposit"),
        ("pink", "move-in date"),
    ]
    # Applied: Format > Highlight > Turquoise's own call, on "Terms".
    control.SetSelection(0, 5)
    wrapper.set_highlight("turquoise")
    assert [p.colour for p in host._highlight_passages()] == ["turquoise", "yellow", "pink"]
    # Removed: Clear Formatting on "deposit".
    start = text.index("deposit")
    control.SetSelection(start, start + len("deposit"))
    wrapper.clear_formatting()
    assert [p.text for p in host._highlight_passages()] == ["Terms", "move-in date"]


def test_rich_offsets_follow_an_edit_above(rich_control) -> None:
    control, wrapper = rich_control
    _load_word_shape(wrapper)
    host = _RichHost(control, wrapper)
    control.SetInsertionPoint(0)
    control.WriteText("Draft. ")
    host.control.SetInsertionPoint(0)
    host.cmd_next_highlight()
    text = control.GetValue()
    assert text[control.GetInsertionPoint() :].startswith("deposit")
    assert host.said[-1] == "Yellow highlight, 1 of 2: deposit"


def test_rich_describe_formatting_now_says_the_highlight(rich_control) -> None:
    control, wrapper = rich_control
    _load_word_shape(wrapper)
    text = control.GetValue()
    control.SetSelection(text.index("deposit") + 1, text.index("deposit") + 1)
    assert wrapper.caret_format_description().endswith("yellow highlight")
    control.SetSelection(text.index("before"), text.index("before"))
    assert "highlight" not in wrapper.caret_format_description()


def test_the_list_dialog_goes_to_the_selected_row_or_nowhere(rich_control) -> None:
    from quill.core.highlights import HighlightPassage
    from quill.ui.highlight_list_dialog import TITLE, show_highlight_list

    passages = [
        HighlightPassage(0, 3, "yellow", "one", "one two"),
        HighlightPassage(8, 11, "pink", "six", "six seven"),
    ]
    outcomes: list[int] = []

    def answer(button_id: int):
        def show(dialog, title):
            assert title == TITLE
            listing = next(c for c in dialog.GetChildren() if isinstance(c, wx.ListBox))
            assert listing.GetSelection() == 1, "opens on the row it was given"
            assert listing.GetString(0) == "one, yellow highlight, in: one two"
            if button_id == wx.ID_OK:
                button = dialog.FindWindowById(wx.ID_OK)
                button.GetEventHandler().ProcessEvent(wx.CommandEvent(wx.wxEVT_BUTTON, wx.ID_OK))
            outcomes.append(button_id)
            return button_id

        return show

    assert show_highlight_list(wx, None, answer(wx.ID_CANCEL), passages, 1) is None
    assert show_highlight_list(wx, None, answer(wx.ID_OK), passages, 1) == 1
    assert outcomes == [wx.ID_CANCEL, wx.ID_OK]


# --------------------------------------------------------------------------- #
# QUILL's adapter
# --------------------------------------------------------------------------- #


def test_quills_adapter_owns_no_command_and_registers_three() -> None:
    from quill.core.keymap import DEFAULT_KEYMAP
    from quill.core.lite.commands import COMMANDS

    assert [n for n in vars(HighlightPassagesMixin) if n.startswith("cmd_")] == []
    lite = {row[3]: row[2] for row in COMMANDS if row[3]}
    for handler, command_id in (
        ("cmd_next_highlight", "navigate.next_highlight"),
        ("cmd_previous_highlight", "navigate.previous_highlight"),
        ("cmd_list_highlights", "navigate.list_highlights"),
    ):
        assert lite[handler] == DEFAULT_KEYMAP[command_id], handler

    registered: list[tuple[str, str]] = []

    class _Commands:
        def try_register(self, command_id, name, handler, binding) -> None:
            registered.append((command_id, binding))

    class _Quill(HighlightPassagesMixin):
        commands = _Commands()

        def _binding_for(self, command_id: str) -> str:
            return DEFAULT_KEYMAP[command_id]

    _Quill().register_highlight_commands()
    assert registered == [
        ("navigate.next_highlight", "Alt+PageDown"),
        ("navigate.previous_highlight", "Alt+PageUp"),
        ("navigate.list_highlights", "Ctrl+F2"),
    ]


def test_quills_host_reads_the_rich_tab_and_jumps_with_back() -> None:
    events: list[Any] = []

    class _Editor:
        def ShowPosition(self, position: int) -> None:  # noqa: N802
            events.append(("show", position))

        def SetFocus(self) -> None:  # noqa: N802
            events.append("focus")

    class _Quill(HighlightPassagesMixin):
        editor = _Editor()

        def _current_editor_mode(self) -> str:
            return "rich"

        def _active_richedit(self) -> str:
            return "wrapper"

        def _record_location_before_jump(self) -> None:
            events.append("recorded")

        def _move_point(self, position: int) -> None:
            events.append(("move", position))

    quill = _Quill()
    assert quill._host_rich_wrapper() == "wrapper"
    quill._host_go_to(12)
    assert events == ["recorded", ("move", 12), ("show", 12), "focus"]
