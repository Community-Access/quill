"""Save As in QUILL writes the format the typed name asks for (2026-10-16, item 41).

Driven through ``MainFrame.save_file_as`` itself, with a stand-in Save As
window and the real writers, and the files checked in their real formats: a
``.docx`` must open in python-docx (a zip with ``word/document.xml``), not be
Markdown text wearing a Word name. Nothing here shows a window.
"""

from __future__ import annotations

import zipfile
from pathlib import Path
from types import SimpleNamespace
from typing import Any

import pytest

from quill.core.document import Document
from quill.ui.main_frame import MainFrame

TEXT = "# Minutes\n\nWe agreed the **budget**.\n\n- Tea\n- Biscuits\n"

ID_OK, ID_CANCEL, YES, NO = 5100, 5101, 5103, 5104


class _Dialog:
    """The Save As window: answers with what the person typed and chose."""

    answer: tuple[int, str, int] = (ID_OK, "", -1)
    opened_on: list[int] = []

    def __init__(self, *_args: Any, **_kwargs: Any) -> None:
        self.filter_index = -1

    def __enter__(self) -> _Dialog:
        return self

    def __exit__(self, *_exc: object) -> None:
        return None

    def SetFilterIndex(self, index: int) -> None:  # noqa: N802 - wx's name
        _Dialog.opened_on.append(index)

    def GetFilterIndex(self) -> int:  # noqa: N802
        return _Dialog.answer[2]

    def GetPath(self) -> str:  # noqa: N802
        return _Dialog.answer[1]


def _wx() -> Any:
    return SimpleNamespace(
        FileDialog=_Dialog,
        FD_SAVE=1,
        FD_OVERWRITE_PROMPT=2,
        ID_OK=ID_OK,
        ID_CANCEL=ID_CANCEL,
        YES=YES,
        NO=NO,
        YES_NO=8,
        YES_DEFAULT=0,
        OK=4,
        ICON_QUESTION=16,
        ICON_INFORMATION=32,
    )


class _Editor:
    def __init__(self, text: str) -> None:
        self.text = text

    def GetValue(self) -> str:  # noqa: N802
        return self.text

    def GetInsertionPoint(self) -> int:  # noqa: N802
        return 0


def _frame(tmp_path: Path, *, export_answer: int = NO) -> MainFrame:
    frame = MainFrame.__new__(MainFrame)
    frame._wx = _wx()
    frame.frame = None
    frame.document = Document(text=TEXT, path=None)
    frame.editor = _Editor(TEXT)
    frame.settings = SimpleNamespace(
        spell_check_before_save=False,
        plain_text_link_style="text_url",
        docx_write_engine="auto",
        default_new_document_format="markdown",
        first_line_as_title=False,
    )
    frame.status = []  # type: ignore[attr-defined]
    frame.boxes = []  # type: ignore[attr-defined]
    frame.exports = []  # type: ignore[attr-defined]
    frame._set_status = frame.status.append  # type: ignore[attr-defined,method-assign]
    frame._announce = lambda _text: None  # type: ignore[method-assign]
    frame._show_modal_dialog = lambda _d, _label: _Dialog.answer[0]  # type: ignore[method-assign]

    def _box(message: str, _caption: str, _style: int) -> int:
        frame.boxes.append(message)  # type: ignore[attr-defined]
        return export_answer

    frame._show_message_box = _box  # type: ignore[method-assign]
    frame.export_document = frame.exports.append  # type: ignore[attr-defined,method-assign]
    frame._file_dialog_default_dir = lambda: str(tmp_path)  # type: ignore[method-assign]
    frame._pinned_markup_kind = lambda: None  # type: ignore[method-assign]
    frame._current_editor_mode = lambda: "markup"  # type: ignore[method-assign]
    frame._save_rich_document_natively = lambda _doc, _target: False  # type: ignore[method-assign]
    for name in (
        "_sync_publishing_linkage_for_document",
        "_record_save_restore_point",
        "_remember_disk_baseline_for",
        "_backup_before_save",
        "_record_recent",
    ):
        setattr(frame, name, lambda *_a, **_k: None)
    for name in (
        "_stop_external_change_watcher",
        "_start_external_change_watcher",
        "_refresh_title",
        "_refresh_sessions_menu",
    ):
        setattr(frame, name, lambda: None)
    frame._load_persistent_undo_state = lambda *_a: None  # type: ignore[method-assign]
    frame._doc_memory = SimpleNamespace(
        set_bookmarks=lambda *_a: None,
        set_anchors=lambda *_a: None,
        set_last_position=lambda *_a: None,
    )
    frame._bookmarks, frame._bookmark_anchors, frame._inline_notes = {}, {}, []
    frame._inline_vault = lambda: SimpleNamespace(set_notes=lambda *_a: None)  # type: ignore[method-assign]
    frame._last_file_dir = ""
    return frame


def _save_as(tmp_path: Path, typed: str, filter_index: int = -1, **kw: Any) -> MainFrame:
    _Dialog.answer = (ID_OK, str(tmp_path / typed), filter_index)
    _Dialog.opened_on = []
    frame = _frame(tmp_path, **kw)
    frame.save_file_as()
    return frame


def _is_word(path: Path) -> bool:
    if not zipfile.is_zipfile(path):
        return False
    with zipfile.ZipFile(path) as archive:
        return "word/document.xml" in archive.namelist()


def test_untitled_text_saved_as_docx_is_a_real_word_document(tmp_path: Path) -> None:
    pytest.importorskip("docx")
    import docx

    _save_as(tmp_path, "minutes.docx")
    target = tmp_path / "minutes.docx"
    assert _is_word(target), "a Word file, not Markdown renamed"
    paragraphs = [p.text for p in docx.Document(str(target)).paragraphs]
    assert "Minutes" in paragraphs
    assert any("budget" in text for text in paragraphs)
    assert "**" not in "".join(paragraphs), "the Markdown became formatting"


def test_an_uppercase_extension_is_the_same_format(tmp_path: Path) -> None:
    pytest.importorskip("docx")
    _save_as(tmp_path, "MINUTES.DOCX")
    assert _is_word(tmp_path / "MINUTES.DOCX")
    _save_as(tmp_path, "NOTES.TXT")
    assert (tmp_path / "NOTES.TXT").read_text(encoding="utf-8") == TEXT


def test_text_and_markdown_are_written_as_they_stand(tmp_path: Path) -> None:
    _save_as(tmp_path, "notes.txt")
    _save_as(tmp_path, "notes.md")
    assert (tmp_path / "notes.txt").read_text(encoding="utf-8") == TEXT
    assert (tmp_path / "notes.md").read_text(encoding="utf-8") == TEXT


def test_html_is_a_page_and_rtf_is_rich_text(tmp_path: Path) -> None:
    _save_as(tmp_path, "page.html")
    page = (tmp_path / "page.html").read_text(encoding="utf-8")
    assert "<html" in page.lower() and "<h1" in page.lower() and "<strong>budget</strong>" in page
    _save_as(tmp_path, "notes.rtf")
    rtf = (tmp_path / "notes.rtf").read_bytes()
    assert rtf.startswith(b"{\\rtf1")
    assert b"**" not in rtf


def test_a_typed_extension_beats_a_mismatched_type_list(tmp_path: Path) -> None:
    # "Word Document" (row 4) highlighted, ".txt" typed: the name wins.
    _save_as(tmp_path, "notes.txt", filter_index=4)
    assert (tmp_path / "notes.txt").read_text(encoding="utf-8") == TEXT
    assert not (tmp_path / "notes.txt.docx").exists()


def test_no_extension_takes_the_highlighted_type(tmp_path: Path) -> None:
    pytest.importorskip("docx")
    _save_as(tmp_path, "minutes", filter_index=4)
    assert _is_word(tmp_path / "minutes.docx")


def test_cancel_writes_nothing(tmp_path: Path) -> None:
    _Dialog.answer = (ID_CANCEL, str(tmp_path / "minutes.docx"), 4)
    frame = _frame(tmp_path)
    frame.save_file_as()
    assert list(tmp_path.iterdir()) == []
    assert frame.document.path is None


def test_an_unsupported_format_is_refused_and_nothing_is_written(tmp_path: Path) -> None:
    frame = _save_as(tmp_path, "budget.xlsx")
    assert not (tmp_path / "budget.xlsx").exists()
    assert frame.boxes and "cannot save to .xlsx" in frame.boxes[0]


def test_pdf_is_offered_to_export_and_never_written_as_text(tmp_path: Path) -> None:
    declined = _save_as(tmp_path, "minutes.pdf", export_answer=NO)
    assert not (tmp_path / "minutes.pdf").exists()
    assert declined.exports == []
    accepted = _save_as(tmp_path, "minutes.PDF", export_answer=YES)
    assert accepted.exports == ["pdf"]
    assert not (tmp_path / "minutes.PDF").exists()
